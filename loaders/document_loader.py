"""
loaders/document_loader.py
──────────────────────────
Responsável por carregar documentos dos formatos suportados.
Retorna uma lista de objetos Document (LangChain) com page_content e metadata.

Formatos suportados:
  • PDF  — via PyPDFLoader
  • TXT  — leitura direta com encoding UTF-8
  • DOCX — via python-docx + wrapper Document
"""

import logging
from pathlib import Path
from typing import List

from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document

logger = logging.getLogger(__name__)


def load_pdf(file_path: str) -> List[Document]:
    """
    Carrega um arquivo PDF e retorna uma lista de Documents,
    um por página, preservando metadados (página, fonte).
    """
    logger.info(f"Carregando PDF: {file_path}")
    loader = PyPDFLoader(file_path)
    docs = loader.load()
    # Garante que o campo 'source' está no metadata
    for doc in docs:
        doc.metadata.setdefault("source", Path(file_path).name)
    logger.info(f"  → {len(docs)} página(s) carregada(s)")
    return docs


def load_txt(file_path: str) -> List[Document]:
    """
    Carrega um arquivo de texto simples.
    Retorna um único Document com todo o conteúdo.
    """
    logger.info(f"Carregando TXT: {file_path}")
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        content = f.read()

    doc = Document(
        page_content=content,
        metadata={"source": Path(file_path).name, "page": 0}
    )
    logger.info("  → 1 documento carregado")
    return [doc]


def load_docx(file_path: str) -> List[Document]:
    """
    Carrega um arquivo DOCX via python-docx.
    Extrai parágrafos não-vazios e os une em um único Document.
    """
    logger.info(f"Carregando DOCX: {file_path}")
    try:
        import docx  # python-docx
    except ImportError:
        raise ImportError("Instale python-docx: pip install python-docx")

    doc_word = docx.Document(file_path)
    paragraphs = [
        para.text.strip()
        for para in doc_word.paragraphs
        if para.text.strip()
    ]
    content = "\n\n".join(paragraphs)

    doc = Document(
        page_content=content,
        metadata={"source": Path(file_path).name, "page": 0}
    )
    logger.info("  → 1 documento carregado")
    return [doc]


def load_document(file_path: str) -> List[Document]:
    """
    Dispatcher: detecta a extensão do arquivo e chama o loader correto.

    Args:
        file_path: Caminho absoluto ou relativo para o arquivo.

    Returns:
        Lista de Documents do LangChain.

    Raises:
        ValueError: Se a extensão não for suportada.
        FileNotFoundError: Se o arquivo não existir.
    """
    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"Arquivo não encontrado: {file_path}")

    ext = path.suffix.lower()

    if ext == ".pdf":
        return load_pdf(file_path)
    elif ext == ".txt":
        return load_txt(file_path)
    elif ext == ".docx":
        return load_docx(file_path)
    else:
        raise ValueError(
            f"Extensão '{ext}' não suportada. "
            f"Use: .pdf, .txt ou .docx"
        )
